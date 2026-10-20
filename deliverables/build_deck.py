"""Build the presentation and the architecture diagram from results/metrics.json.

python deliverables/build_deck.py  ->  deliverables/GlycoTwin_Presentation.pptx, deliverables/GlycoTwin_Architecture.pptx
"""
import json
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
M = json.load(open(os.path.join(ROOT, "results", "metrics.json")))
OP = json.load(open(os.path.join(ROOT, "results", "operating.json")))
OPH, OPO = OP["hyper"], OP["hypo"]
FULLN = "Full twin (+ EHR)"
FIG = os.path.join(ROOT, "figures")

INK = RGBColor(0x1D, 0x23, 0x21)
MUTE = RGBColor(0x66, 0x70, 0x6C)
GREEN = RGBColor(0x1A, 0x6B, 0x54)
LINE = RGBColor(0xC9, 0xCC, 0xC9)
PALE = RGBColor(0xE7, 0xF1, 0xEC)
ORANGE = RGBColor(0xC2, 0x57, 0x1A)
BLUE = RGBColor(0x2F, 0x5F, 0xB3)
FONT = "Helvetica Neue"


def new_prs():
    p = Presentation()
    p.slide_width, p.slide_height = Inches(13.333), Inches(7.5)
    return p


def text(slide, x, y, w, h, s, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT, font=FONT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    paras = s if isinstance(s, list) else [s]
    for i, ptxt in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        if isinstance(ptxt, tuple):
            ptxt, kw = ptxt
        else:
            kw = {}
        r = para.add_run()
        r.text = ptxt
        r.font.size = Pt(kw.get("size", size))
        r.font.bold = kw.get("bold", bold)
        r.font.name = font
        r.font.color.rgb = kw.get("color", color)
        para.space_after = Pt(kw.get("after", 6))
    return tb


def base(prs, claim, kicker=None, n=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    if kicker:
        text(s, 0.6, 0.35, 12, 0.35, kicker.upper(), size=11, color=GREEN, bold=True)
    text(s, 0.6, 0.62, 12.1, 1.1, claim, size=26, bold=True)
    if n:
        text(s, 12.2, 7.0, 0.8, 0.3, str(n), size=10, color=MUTE, align=PP_ALIGN.RIGHT)
    text(s, 0.6, 7.0, 8, 0.3, "GlycoTwin · Team Snack Overflow, RVCE", size=10, color=MUTE)
    return s


def rule(slide, x, y, w, weight=0.75, color=INK):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + w), Inches(y))
    c.line.color.rgb = color
    c.line.width = Pt(weight)


def _numeric(c):
    import re
    return bool(re.fullmatch(r"[\[\]\d.,%+\-–\s/a-z]*", str(c))) and any(ch.isdigit() for ch in str(c)) or str(c) in ("–", "0")


def table(slide, x, y, colw, rows, header, size=14, rowh=0.42, bold_rows=()):
    """Booktabs-style table: thick top and bottom rules, thin mid rule, no verticals."""
    w = sum(colw)
    num_col = [all(_numeric(r[j]) for r in rows) for j in range(len(header))]
    rule(slide, x, y, w, 1.5)
    cx = x
    for j, hcell in enumerate(header):
        text(slide, cx, y + 0.06, colw[j], rowh, hcell, size=size - 1, bold=True,
             align=PP_ALIGN.RIGHT if (j and num_col[j]) else PP_ALIGN.LEFT, color=MUTE)
        cx += colw[j]
    yy = y + rowh + 0.04
    rule(slide, x, yy, w, 0.5)
    for i, r in enumerate(rows):
        cx = x
        for j, cell in enumerate(r):
            text(slide, cx, yy + 0.04, colw[j], rowh, cell, size=size, bold=i in bold_rows,
                 align=PP_ALIGN.RIGHT if (j and num_col[j]) else PP_ALIGN.LEFT, color=GREEN if i in bold_rows else INK)
            cx += colw[j]
        yy += rowh
    rule(slide, x, yy + 0.06, w, 1.5)
    return yy + 0.1


def box(slide, x, y, w, h, title, body, fill=RGBColor(0xFF, 0xFF, 0xFF), edge=INK, tsize=13, bsize=11):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = 0.08
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = edge
    shp.line.width = Pt(1)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, Inches(0.08))
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = title
    r.font.size, r.font.bold, r.font.name, r.font.color.rgb = Pt(tsize), True, FONT, INK
    for line in body:
        q = tf.add_paragraph()
        q.alignment = PP_ALIGN.LEFT
        rr = q.add_run()
        rr.text = line
        rr.font.size, rr.font.name, rr.font.color.rgb = Pt(bsize), FONT, MUTE
    return shp


def arrow(slide, x1, y1, x2, y2, color=INK):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(1.25)
    ln = c.line._get_or_add_ln()
    from lxml import etree
    tail = etree.SubElement(ln, "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd")
    tail.set("type", "triangle")
    return c


def architecture(s, top=1.7):
    # sources
    box(s, 0.6, top, 2.7, 1.55, "Static stream: EHR",
        ["one row per monitoring period:", "age, sex, BMI, years with T2D,", "complications, comorbidities,",
         "16 labs (HbA1c, C-peptide, eGFR...),", "drugs, past hypoglycaemia"], fill=PALE)
    box(s, 0.6, top + 1.85, 2.7, 1.55, "Dynamic stream: sensor + logs",
        ["CGM every 15 min (mg/dL)", "meals (free text, grams)", "s.c. insulin, pump bolus/basal,",
         "oral agents"], fill=PALE)
    box(s, 0.6, top + 3.7, 2.7, 1.0, "Planned, not in this PoC",
        ["FHIR R4 / ABDM records", "Apple Health, Google Fit: HR, steps, sleep"], edge=LINE)
    # ingestion
    box(s, 3.75, top + 0.6, 2.6, 2.4, "Ingest and align",
        ["regular 15-min grid per patient", "gap fill only up to 30 min", "parse meals (g), insulin (IU)",
         "EHR fields to numbers,", "HOMA-IR derived", "causal: nothing after t"])
    # twin state
    box(s, 6.8, top, 2.9, 3.4, "Twin state at time t",
        ["CGM features: level, lags, trends,", "2 h and 6 h windows, time of day", "",
         "Event features: time since meal, food", "in last 2 h, insulin on board, oral agent",
         "", "Twin memory (learned per patient):", "personal mean, time above range,",
         "post-meal rise, usual meal times", "", "EHR profile: 33 static fields"], fill=PALE, edge=GREEN)
    # models
    box(s, 10.15, top, 2.6, 1.0, "Hyper head", ["P(rise above 180 mg/dL", "within 2 h)"], edge=ORANGE)
    box(s, 10.15, top + 1.15, 2.6, 1.0, "Hypo head", ["P(fall below 70 mg/dL", "within 2 h)"], edge=BLUE)
    box(s, 10.15, top + 2.3, 2.6, 1.1, "Trajectory head", ["glucose at +30, +60,", "+90, +120 min"], edge=GREEN)
    box(s, 6.8, top + 3.7, 5.95, 1.0, "Clinician dashboard",
        ["EHR card · twin memory · CGM with forecast · 2 h risk timeline · alert threshold per ward · top-5 drivers (tree SHAP)"],
        edge=INK)
    arrow(s, 3.3, top + 0.8, 3.75, top + 1.4)
    arrow(s, 3.3, top + 2.6, 3.75, top + 2.2)
    arrow(s, 6.35, top + 1.8, 6.8, top + 1.8)
    for yy in (top + 0.5, top + 1.65, top + 2.85):
        arrow(s, 9.7, top + 1.7, 10.15, yy)
    arrow(s, 11.45, top + 3.4, 11.45, top + 3.7)
    text(s, 10.15, top - 0.32, 2.6, 0.3, "LightGBM, one model per head", size=10, color=MUTE)


def pct(x):
    return f"{100 * x:.0f}%"


def build():
    H, O, F, D = M["hyper"], M["hypo"], M["forecast"], M["dataset"]
    names = ["Trend rule (no ML)", "EHR only", "CGM only", "CGM + events", "CGM + events + twin memory", "Full twin (+ EHR)"]
    prs = new_prs()
    n = 0

    # 1 title
    s = prs.slides.add_slide(prs.slide_layouts[6])
    text(s, 0.8, 1.6, 11.5, 0.5, "HAPPIEST HEALTH DIGITAL TWIN CHALLENGE 2026 · PHASE 1", size=12, color=GREEN, bold=True)
    text(s, 0.8, 2.1, 11.5, 1.0, "GlycoTwin", size=54, bold=True)
    text(s, 0.8, 3.2, 11.5, 1.2, "A type 2 diabetes digital twin that fuses the health record with a glucose sensor\nand warns two hours before glucose leaves the safe range", size=22, color=MUTE)
    rule(s, 0.8, 4.75, 11.7, 0.75, LINE)
    text(s, 0.8, 4.95, 11.5, 1.6, [
        ("Team Snack Overflow · RV College of Engineering, Bengaluru", {"bold": True, "size": 16}),
        ("Yash Saraogi (lead) · Tallam Sri Sai Subramanyam · Yash Gautam", {"size": 15, "color": MUTE}),
        ("github.com/chxmq/SnackOverflow_RVCE", {"size": 15, "color": GREEN})])

    # 2 problem
    n += 1
    s = base(prs, "Between two clinic visits a doctor sees one HbA1c; the patient lives through 8,640 glucose readings", "The problem", n + 1)
    text(s, 0.6, 1.9, 6.2, 4.8, [
        ("101 million Indians live with diabetes and 136 million with prediabetes (ICMR-INDIAB, Lancet Diabetes & Endocrinology, 2023).", {"size": 17, "after": 14}),
        ("HbA1c is a three-month average. It cannot say when glucose went high, or how low it fell at 3 am.", {"size": 17, "after": 14}),
        ("A 15-minute CGM records 96 readings a day: 8,640 per quarter. No clinician can read that raw.", {"size": 17, "after": 14}),
        ("Hypoglycaemia on insulin or sulfonylureas is the acute danger; repeated post-meal spikes are the slow one.", {"size": 17}),
    ])
    box(s, 7.3, 1.95, 5.4, 4.4, "What is missing",
        ["", "A model of this one patient that:", "", "  1. reads the record AND the live stream",
         "  2. learns how this person responds to meals", "  3. says, two hours early, which way",
         "      glucose is about to leave range", "  4. shows a doctor why it thinks so", "",
         "That is our digital twin, scoped to one", "condition: type 2 diabetes."], fill=PALE, edge=GREEN, tsize=16, bsize=15)

    # 3 use case
    n += 1
    s = base(prs, "Use case: a diabetes team watches a panel of CGM-wearing patients and acts before the excursion, not after", "Healthcare use case", n + 1)
    table(s, 0.6, 1.95, [3.0, 4.6, 4.5], [
        ["Who", "Endocrinologist, diabetes educator, ward nurse", "Inpatient ward or remote monitoring"],
        ["Input", "Record at admission + live CGM + meal and dose logs", "No new hardware beyond a CGM"],
        ["Hyper alert", "Rise above 180 mg/dL within 2 h", "Check meal, correction dose, walk"],
        ["Hypo alert", "Fall below 70 mg/dL within 2 h", "Snack, hold sulfonylurea / insulin"],
        ["Twin view", "Personal post-meal rise, usual meal times", "Tailor dose and diet advice"],
        ["Explanation", "Top 5 drivers per alert", "Doctor can overrule a wrong reason"],
    ], ["", "What the twin does", "What the clinician can do"], size=15, rowh=0.55)
    text(s, 0.6, 6.0, 12, 0.8, "The twin does not dose or prescribe. It is decision support: it ranks who needs attention in the next two hours and explains why.", size=15, color=MUTE)

    # 4 data
    n += 1
    s = base(prs, f"We used real, anonymised patient data: {D['patients']} people with type 2 diabetes, {D['patient_days']:.0f} patient-days of CGM", "Data and sandbox rules", n + 1)
    table(s, 0.6, 1.95, [3.2, 2.0, 6.9], [
        ["Dataset", "", "ShanghaiT2DM (Zhao et al., Scientific Data 10:35, 2023), CC BY 4.0"],
        ["Patients / periods", f"{D['patients']} / {D['records']}", "Some patients were monitored in more than one period"],
        ["CGM points", f"{D['cgm_points']:,}", "15-min interval, 3 to 14 days per period"],
        ["Static (EHR)", "33 fields", "Age, sex, BMI, smoking, diabetes years, complications, comorbidities, drugs"],
        ["Labs", "16 tests", "FPG, 2 h PPG, C-peptide, insulin, HbA1c, glycated albumin, lipids, eGFR, uric acid"],
        ["Dynamic logs", "4 kinds", "Meals with grams, s.c. insulin, pump bolus and basal, oral agents"],
    ], ["", "Size", "Detail"], size=14, rowh=0.5)
    text(s, 0.6, 5.65, 12.1, 1.2, [
        ("Why not only synthetic data: a model trained on our own simulator would learn our simulator. Real CGM keeps us honest.", {"size": 14}),
        ("Sandbox rules: the data is public, de-identified and openly licensed. No identifiable patient data is used anywhere (DPDP Act and HIPAA).", {"size": 14, "color": MUTE}),
    ])

    # 5 architecture
    n += 1
    s = base(prs, "Architecture: two streams in, one twin state every 15 minutes, three predictions out", "System design", n + 1)
    architecture(s)

    # 6 twin definition
    n += 1
    s = base(prs, "What makes it a twin and not just a forecaster: it keeps a memory of this patient and updates it as data arrives", "The twin", n + 1)
    table(s, 0.6, 1.95, [3.6, 5.2, 3.3], [
        ["Personal mean glucose", "Running mean of every CGM reading so far", "starts after 2 h"],
        ["Personal time above range", "Share of readings ≥ 180 mg/dL so far", "starts after 2 h"],
        ["Personal post-meal rise", "Mean peak rise in the 2 h after each past meal", "after first meal + 2 h"],
        ["Usual meal in next 2 h", "Share of past days with a meal logged in this clock window", "from day 2"],
        ["EHR profile", "HbA1c, C-peptide, HOMA-IR, drugs, complications", "fixed per period"],
    ], ["Twin memory item", "How it is learned", "Available"], size=14, rowh=0.55)
    text(s, 0.6, 5.3, 12.1, 1.5, [
        ("Each item is computed only from data before time t. A meal counts toward the post-meal rise only once its 2-hour window has fully passed, so the label never leaks into the feature.", {"size": 14}),
        ("The twin replays forward: drag the dashboard cursor and the memory grows exactly as it would have at the bedside.", {"size": 14, "color": MUTE}),
    ])

    # 7 protocol
    n += 1
    s = base(prs, "Evaluation protocol: every number in this deck is on patients the model never saw", "Method", n + 1)
    text(s, 0.6, 1.9, 6.0, 4.8, [
        ("Targets", {"bold": True, "size": 16}),
        ("Hyper: glucose now < 180, and any reading ≥ 180 in the next 2 h.", {"size": 15}),
        ("Hypo: glucose now ≥ 70, and any reading < 70 in the next 2 h.", {"size": 15, "after": 14}),
        ("Splits", {"bold": True, "size": 16}),
        ("5-fold GroupKFold by patient. A patient's periods never straddle train and test.", {"size": 15, "after": 14}),
        ("Alert threshold", {"bold": True, "size": 16}),
        ("Fixed at p ≥ 0.5 before looking at any test fold. Not tuned.", {"size": 15}),
    ])
    text(s, 7.0, 1.9, 5.7, 4.8, [
        ("Uncertainty", {"bold": True, "size": 16}),
        ("95% confidence intervals by bootstrapping patients (500 resamples), not windows, because windows from one patient are correlated.", {"size": 15, "after": 14}),
        ("Baselines", {"bold": True, "size": 16}),
        ("Trend rule: current glucose plus twice the last 30-min change, what a nurse would eyeball.", {"size": 15}),
        ("Ablations: EHR only, CGM only, CGM + events, + twin memory, + EHR.", {"size": 15, "after": 14}),
        ("Windows", {"bold": True, "size": 16}),
        (f"Hyper: {H['n_windows']:,} windows, {pct(H['prevalence'])} positive. Hypo: {O['n_windows']:,} windows, {100 * O['prevalence']:.1f}% positive.", {"size": 15}),
    ])

    # 8 results table
    n += 1
    PT, CE = "Personal twin (CGM + events + memory)", "Personal twin + core EHR (9 fields)"
    s = base(prs, f"On patients it never saw, the twin flags highs at AUROC {OPH[PT]['auroc']:.3f} and lows at {OPO[PT]['auroc']:.3f}, well above a nurse's trend rule", "Results", n + 1)
    rows = []
    for nm in names[:4]:
        rows.append([nm, f"{H[nm]['auroc']:.3f}", f"{H[nm]['auprc']:.3f}", f"{O[nm]['auroc']:.3f}", f"{O[nm]['auprc']:.3f}"])
    rows.append(["Personal twin: CGM + events + twin memory", f"{OPH[PT]['auroc']:.3f}", f"{OPH[PT]['auprc']:.3f}", f"{OPO[PT]['auroc']:.3f}", f"{OPO[PT]['auprc']:.3f}"])
    rows.append(["Personal twin + 9 core EHR fields", f"{OPH[CE]['auroc']:.3f}", f"{OPH[CE]['auprc']:.3f}", f"{OPO[CE]['auroc']:.3f}", f"{OPO[CE]['auprc']:.3f}"])
    rows.append(["Full twin: + all 33 EHR fields (dashboard)", f"{H[FULLN]['auroc']:.3f}", f"{H[FULLN]['auprc']:.3f}", f"{O[FULLN]['auroc']:.3f}", f"{O[FULLN]['auprc']:.3f}"])
    table(s, 0.6, 1.85, [5.0, 1.75, 1.75, 1.75, 1.75], rows,
          ["Model (5-fold, patient-held-out)", "Hyper AUROC", "Hyper AUPRC", "Hypo AUROC", "Hypo AUPRC"], size=14, rowh=0.44, bold_rows=(4,))
    lo, hi = H["full_minus_rule_auroc_ci95"]
    lo2, hi2 = O["full_minus_rule_auroc_ci95"]
    text(s, 0.6, 5.75, 12.1, 1.2, [
        (f"Full twin minus trend rule, AUROC, 95% CI over patients: highs [{lo:+.3f}, {hi:+.3f}], lows [{lo2:+.3f}, {hi2:+.3f}].", {"size": 14}),
        (f"Chance level for AUPRC is the prevalence: {H['prevalence']:.3f} for highs, {O['prevalence']:.3f} for lows.", {"size": 13, "color": MUTE}),
    ])

    # 9 figure
    n += 1
    s = base(prs, "The timing comes from the live sensor, the twin's memory of the patient adds to it, and the record alone cannot time an event", "Ablation", n + 1)
    s.shapes.add_picture(os.path.join(FIG, "ablation_auroc.png"), Inches(0.6), Inches(1.95), width=Inches(12.1))

    # 9b fusion, honestly
    n += 1
    s = base(prs, "Honest finding: once the twin has a few hours of glucose, the static record adds no measurable accuracy", "What fusion with the record bought us", n + 1)
    ci = lambda r, k: f"[{r['ci_vs_cgm'][k][0]:+.3f}, {r['ci_vs_cgm'][k][1]:+.3f}]"
    d1 = lambda r, k: "–" if r[k]["auroc_day1"] is None else f"{r[k]['auroc_day1']:.3f}"
    table(s, 0.6, 1.85, [4.6, 2.2, 2.2, 1.55, 1.55], [
        ["Personal twin (memory, no EHR)", ci(OPH, PT), ci(OPO, PT), d1(OPH, PT), d1(OPO, PT)],
        ["+ 9 core EHR fields", ci(OPH, CE), ci(OPO, CE), d1(OPH, CE), d1(OPO, CE)],
        ["+ all 33 EHR fields", ci(OPH, FULLN), ci(OPO, FULLN), d1(OPH, FULLN), d1(OPO, FULLN)],
        ["CGM only (reference)", "0", "0", d1(OPH, "CGM only"), d1(OPO, "CGM only")],
    ], ["AUROC gain over CGM only, 95% CI", "Highs", "Lows", "Day-1 highs", "Day-1 lows"], size=14, rowh=0.5)
    text(s, 0.6, 4.6, 12.1, 2.3, [
        ("Highs: the memory gain excludes zero. Lows: every interval includes zero, so we claim no gain there.", {"size": 14, "after": 10, "bold": True}),
        ("Why the record adds nothing here: the twin's learned mean glucose and time above range are a current, personal version of what HbA1c summarises over three months. With 100 patients, 33 static fields mostly let trees memorise who is who.", {"size": 14, "after": 10}),
        (f"The record still carries population risk on its own (EHR-only AUROC {H['EHR only']['auroc']:.3f} for highs), and the clinician needs it beside every alert. In the dashboard it is context, not the source of the timing.", {"size": 14, "after": 10}),
        ("Day 1 = the first calendar day of monitoring, before the memory has filled. Even there the record did not add accuracy in this cohort.", {"size": 13, "color": MUTE}),
    ])

    # 10 event level
    n += 1
    tp = {r["threshold"]: r for r in OPH[PT]["tradeoff"]}
    to = {r["threshold"]: r for r in OPO[PT]["tradeoff"]}
    s = base(prs, f"What a ward would see: at p ≥ 0.5 the twin catches {pct(tp[0.5]['recall'])} of highs before onset; lows are rarer and need a lower threshold", "Alert trade-off", n + 1)
    s.shapes.add_picture(os.path.join(FIG, "tradeoff.png"), Inches(0.6), Inches(1.8), width=Inches(8.2))
    def lead(e):
        return "–" if e["median_lead_min"] is None else f"{e['median_lead_min']:.0f}"
    table(s, 9.0, 1.9, [1.3, 1.0, 0.9, 0.95], [
        ["Highs 0.5", pct(tp[0.5]["recall"]), lead(tp[0.5]), f"{tp[0.5]['false_per_day']:.2f}"],
        ["Highs 0.3", pct(tp[0.3]["recall"]), lead(tp[0.3]), f"{tp[0.3]['false_per_day']:.2f}"],
        ["Lows 0.5", pct(to[0.5]["recall"]), lead(to[0.5]), f"{to[0.5]['false_per_day']:.2f}"],
        ["Lows 0.2", pct(to[0.2]["recall"]), lead(to[0.2]), f"{to[0.2]['false_per_day']:.2f}"],
        ["Lows 0.1", pct(to[0.1]["recall"]), lead(to[0.1]), f"{to[0.1]['false_per_day']:.2f}"],
    ], ["Personal twin", "Caught", "Lead min", "False/day"], size=12, rowh=0.4)
    text(s, 9.0, 4.65, 3.8, 2.3, [
        (f"Excursions: {tp[0.5]['events']:,} highs, {to[0.5]['events']:,} lows in held-out data. Caught = an alert in the 2 h before the crossing. False = an alert run with no crossing in the next 2 h.", {"size": 11}),
        ("0.5 was fixed in advance. The other rows are read off the same curve to show the cost, not tuned results.", {"size": 11, "color": MUTE}),
    ])

    # 11 forecast + calibration
    n += 1
    s = base(prs, f"The twin also forecasts the trajectory: {F['60']['mae']:.1f} mg/dL mean error at 60 min versus {F['60']['persistence_mae']:.1f} for 'no change'", "Trajectory and calibration", n + 1)
    s.shapes.add_picture(os.path.join(FIG, "forecast_mae.png"), Inches(0.6), Inches(1.9), height=Inches(4.3))
    s.shapes.add_picture(os.path.join(FIG, "calibration.png"), Inches(6.6), Inches(1.9), height=Inches(3.0))
    text(s, 6.6, 5.0, 6.1, 1.6, [
        ("Calibration (right): predicted risk against the observed rate in held-out patients. Points near the diagonal mean a 30% alert really means about 30%.", {"size": 13}),
        (f"Per-patient AUROC for the hyper head: median {H['per_patient_auroc']['median']:.2f} (IQR {H['per_patient_auroc']['q25']:.2f} to {H['per_patient_auroc']['q75']:.2f}) across {H['per_patient_auroc']['n']} patients with events.", {"size": 13, "color": MUTE}),
    ])

    # 12 dashboard
    n += 1
    s = base(prs, "The doctor's view: one screen per virtual patient, replayable, with the reason for every alert", "Clinician dashboard", n + 1)
    shot = os.path.join(FIG, "dashboard.png")
    if os.path.exists(shot):
        s.shapes.add_picture(shot, Inches(1.7), Inches(1.7), height=Inches(4.75))
    text(s, 0.6, 6.55, 12.1, 0.4, "Live: chxmq.github.io/SnackOverflow_RVCE  ·  every number shown is an out-of-fold prediction for a held-out patient", size=12, color=MUTE)

    # 13 drivers
    n += 1
    from glycotwin.train import PRETTY
    def merged(d):
        acc = {}
        for k, v in d.items():
            acc[PRETTY.get(k, k)] = acc.get(PRETTY.get(k, k), 0) + v
        return sorted(acc.items(), key=lambda kv: -kv[1])[:8]
    top, topo = merged(H["top_features"]), merged(O["top_features"])
    s = base(prs, "What drives the twin: glucose now and its trend, the clock, and what the twin has learned about this patient", "Explainability", n + 1)
    table(s, 0.6, 1.95, [4.2, 1.6], [[k, f"{v:.3f}"] for k, v in top], ["Hyper head: feature", "mean |SHAP|"], size=13, rowh=0.42)
    table(s, 7.0, 1.95, [4.2, 1.6], [[k, f"{v:.3f}"] for k, v in topo], ["Hypo head: feature", "mean |SHAP|"], size=13, rowh=0.42)
    text(s, 0.6, 5.85, 12.1, 1.0, [("Time of day ranks high for highs because this is a hospital cohort with fixed meal times. In outpatients we expect logged meals to take that role.", {"size": 13}),
                                    ("Mean absolute tree-SHAP contribution in log-odds on held-out predictions (full twin). The dashboard shows the same values per alert.", {"size": 12, "color": MUTE})])

    # 14 limitations
    n += 1
    s = base(prs, "What this proof of concept does not show yet", "Limitations, stated plainly", n + 1)
    text(s, 0.6, 1.9, 12.1, 5.0, [
        ("One cohort: Chinese adults, mostly monitored in hospital. Indian diets, outpatient life and Indian CGM users need their own validation.", {"size": 16, "after": 12}),
        ("Short memory: 3 to 14 days per period, so the twin learns habits from at most two weeks.", {"size": 16, "after": 12}),
        ("No heart rate, steps or sleep in this dataset. The wearable adapters are designed but untested, so we make no claim about them.", {"size": 16, "after": 12}),
        ("Meal size comes from logged grams of food, not carbohydrate counts.", {"size": 16, "after": 12}),
        ("Retrospective only. No prospective or silent-mode trial yet, and not a medical device.", {"size": 16, "after": 12}),
        ("Associations, not causes: the twin warns and explains, but it does not simulate the effect of a dose change.", {"size": 16}),
    ])

    # 15 roadmap
    n += 1
    s = base(prs, "Path to an Indian deployment", "Next steps", n + 1)
    table(s, 0.6, 1.95, [3.0, 6.1, 3.0], [
        ["Wearables", "Add HR, HRV, steps and sleep from Apple Health / Google Fit as dynamic features", "Measure the AUROC gain"],
        ["Records", "Read EHR through FHIR R4 / ABDM instead of a spreadsheet", "Same twin, live input"],
        ["Indian cohort", "Retrain and recalibrate on Indian CGM data with consent under DPDP", "Calibration per site"],
        ["Silent trial", "Run beside a ward for 4 weeks; compare alerts with what happened", "Alert burden, lead time"],
        ["Mechanistic core", "Pair with a physiological glucose model for what-if dose simulation", "True counterfactuals"],
    ], ["Step", "What", "Success measure"], size=14, rowh=0.62)

    # 16 close
    n += 1
    s = base(prs, "Summary", None, n + 1)
    text(s, 0.6, 1.8, 12.1, 4.0, [
        (f"Real data: {D['patients']} patients with type 2 diabetes, EHR plus 15-min CGM, fused into one twin state.", {"size": 18, "after": 12}),
        (f"Held-out patients: AUROC {OPH['Personal twin (CGM + events + memory)']['auroc']:.3f} for highs and {OPO['Personal twin (CGM + events + memory)']['auroc']:.3f} for lows two hours ahead, clearly above a trend rule.", {"size": 18, "after": 12}),
        ("Measured, not assumed: the twin's memory of the patient adds accuracy for highs; the static record, in this cohort, does not.", {"size": 18, "after": 12}),
        ("A clinician dashboard that replays the twin and explains every alert.", {"size": 18, "after": 12}),
        ("Open source under MIT. One command reproduces every number.", {"size": 18}),
    ])
    rule(s, 0.6, 5.6, 12.1, 0.75, LINE)
    text(s, 0.6, 5.75, 12.1, 1.0, [("github.com/chxmq/SnackOverflow_RVCE  ·  chxmq.github.io/SnackOverflow_RVCE", {"size": 16, "color": GREEN}),
                                    ("Team Snack Overflow · RV College of Engineering", {"size": 14, "color": MUTE})])
    prs.save(os.path.join(HERE, "GlycoTwin_Presentation.pptx"))

    a = new_prs()
    s = a.slides.add_slide(a.slide_layouts[6])
    text(s, 0.6, 0.35, 12, 0.35, "GLYCOTWIN · ARCHITECTURE", size=11, color=GREEN, bold=True)
    text(s, 0.6, 0.62, 12.1, 1.0, "Two data streams in, one twin state every 15 minutes, three predictions out", size=24, bold=True)
    architecture(s)
    text(s, 0.6, 7.0, 12, 0.3, "Team Snack Overflow, RVCE · github.com/chxmq/SnackOverflow_RVCE", size=10, color=MUTE)
    a.save(os.path.join(HERE, "GlycoTwin_Architecture.pptx"))


if __name__ == "__main__":
    build()
